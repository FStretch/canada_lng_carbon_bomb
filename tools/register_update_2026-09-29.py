"""Move LNG Canada Phase 2 to under construction after its final investment decision.

The LNG Canada partners took the decision on 29 September 2026 (announced
28 September Pacific time). Under the paper's rule, committed means past a
final investment decision, so the row leaves advanced_proposed.

The script checks the current row before it writes anything, and stops if
those values are not the ones this update expects.

Usage: python register_update_2026-09-29.py path/to/Canada_LNG_Asset_Register.xlsx
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

import openpyxl

PID = "lng_canada_phase_2"

LNGC = (
    "LNG Canada, LNG Canada Announces Phase 2 Final Investment Decision (Sep 2026): "
    "https://www.newswire.ca/news-releases/"
    "lng-canada-announces-phase-2-final-investment-decision-828313121.html"
)
SHELL = (
    "Shell, Shell takes final investment decision on LNG Canada Phase 2 (29 Sep 2026): "
    "https://news.europawire.eu/shell-takes-final-investment-decision-on-lng-canada-phase-2-"
    "to-double-kitimat-production-capacity-to-28-million-tonnes-per-year/eu-press-release/"
    "2026/09/29/14/54/00/181489/"
)
FT = (
    "Financial Times, Shell-led consortium backs $23bn expansion of LNG Canada project "
    "(29 Sep 2026): https://www.ft.com/content/364d5454-876d-42ef-8f30-759e1ebdb026"
)
FLUOR = (
    "Fluor, Fluor joint venture selected for LNG Canada Phase 2 expansion following "
    "final investment decision (29 Sep 2026): "
    "https://www.financialcontent.com/article/bizwire-2026-9-29-fluor-joint-venture-selected-"
    "for-lng-canada-phase-2-expansion-following-final-investment-decision"
)
TC = (
    "TC Energy, Coastal GasLink Phase 2 to proceed following LNG Canada Final Investment "
    "Decision (29 Sep 2026): "
    "https://www.globenewswire.com/news-release/2026/09/29/3370625/0/en/"
    "coastal-gaslink-phase-2-to-proceed-following-lng-canada-final-investment-decision.html"
)
NEW_SOURCES = (LNGC, SHELL, FT, FLUOR, TC)

TIER_REASON = (
    "Final investment decision taken 29 September 2026 by the LNG Canada joint venture "
    "(Shell, PETRONAS, PetroChina, Mitsubishi, KOGAS); Fluor joint venture selected for "
    "engineering, procurement and construction. Moved from advanced_proposed."
)
COMMERCIAL_APPEND = (
    "FID 29 Sep 2026; phase 2 cost reported at US$21\u201323bn (Mitsubishi share "
    "US$3.2bn implies US$21.3bn); total terminal capacity 28 mtpa; commercial "
    "operation expected in the early 2030s."
)
FIRST_EXPORT_SOURCE = (
    "2031: decision date September 2026 plus the IEA's 4 to 5 year lag from final "
    "investment decision to completion (Parameters post_fid_to_completion = 4.5), "
    "consistent with Shell's statement that commercial operations are expected in "
    "the early 2030s. Decision 29 Sep 2026. "
    + SHELL
)
CGL_NOTE = (
    "Phase 2 expansion to proceed following the LNG Canada Phase 2 FID (29 Sep 2026)"
)

# Headline export capacity after Phase 2 leaves proposed. Checked against the
# register after the row is written, not assumed in place of that sum.
EXPECTED_EXPORT = {
    "operating": 14.0,
    "under_construction": 19.85,
    "proposed": 51.7,
}


def fail(msg: str) -> None:
    raise SystemExit(f"ASSERT FAILED: {msg}")


class Register:
    def __init__(self, path: str):
        self.path = path
        self.wb = openpyxl.load_workbook(path)
        self.reg = self.wb["Asset Register"]
        self.src = self.wb["Asset Sources"]
        self.cols = {c.value: c.column for c in self.reg[1] if c.value}
        self.scols = {c.value: c.column for c in self.src[1] if c.value}
        if set(self.cols) != set(self.scols):
            fail("Asset Register and Asset Sources fields differ")
        self.rows = {}
        for r in range(2, self.reg.max_row + 1):
            pid = self.reg.cell(r, 1).value
            if pid:
                self.rows[pid] = r
                if self.src.cell(r, 1).value != pid:
                    fail(f"Asset Sources row mismatch for {pid}")
        self.log: list[tuple[str, str, object, object]] = []

    def get(self, pid: str, field: str):
        return self.reg.cell(self.rows[pid], self.cols[field]).value

    def set(self, pid: str, field: str, value, source: str | None = None):
        old = self.get(pid, field)
        self.reg.cell(self.rows[pid], self.cols[field]).value = value
        if source is not None:
            sold = self.src.cell(self.rows[pid], self.scols[field]).value
            self.src.cell(self.rows[pid], self.scols[field]).value = source
            self.log.append((f"Asset Sources {pid}", field, sold, source))
        self.log.append((f"Asset Register {pid}", field, old, value))

    def add_sources(self, pid: str, *entries: str):
        cur = self.get(pid, "sources") or ""
        parts = [p.strip() for p in str(cur).split("|") if p.strip()]
        for entry in entries:
            url = entry.split(": ")[-1].strip()
            if not any(url in p for p in parts):
                parts.append(entry)
        compiled = " | ".join(parts)
        # The Asset Sources "sources" cell is the boilerplate "Compiled by us...",
        # not a second copy of the register list. Field cells above carry the
        # new citations; this column on the register is the compiled list.
        old = self.get(pid, "sources")
        self.reg.cell(self.rows[pid], self.cols["sources"]).value = compiled
        self.log.append((f"Asset Register {pid}", "sources", old, compiled))


def _blank(val) -> bool:
    if val is None:
        return True
    if isinstance(val, str) and val.strip() == "":
        return True
    return False


def _fid_is_false(val) -> bool:
    return val is False or val in (0, 0.0)


def assert_current(R: Register) -> None:
    if PID not in R.rows:
        fail(f"{PID} is not on the Asset Register")
    fid = R.get(PID, "fid_confirmed")
    tier = R.get(PID, "tier")
    calc = R.get(PID, "calc_group")
    year = R.get(PID, "first_export_year")
    if not _fid_is_false(fid):
        fail(f"fid_confirmed is {fid!r}, expected False")
    if tier != "advanced_proposed":
        fail(f"tier is {tier!r}, expected advanced_proposed")
    if calc != "proposed":
        fail(f"calc_group is {calc!r}, expected proposed")
    if int(year) != 2030:
        fail(f"first_export_year is {year!r}, expected 2030")
    status = R.get(PID, "status")
    if status != "proposed":
        fail(f"status is {status!r}; Cedar and Woodfibre use construction after FID")
    for other in ("cedar_lng", "woodfibre_lng"):
        if R.get(other, "status") != "construction":
            fail(f"{other} status is {R.get(other, 'status')!r}, expected construction")
    fd = R.wb["Field Definitions"]
    allowed = None
    for row in fd.iter_rows(min_row=2, max_col=3, values_only=True):
        if row[0] == "status" and isinstance(row[2], str):
            allowed = row[2]
    if allowed is None or "construction" not in allowed:
        fail(f"Field Definitions status text does not list construction: {allowed!r}")


def update_phase_2(R: Register) -> None:
    note = R.get(PID, "commercial_note")
    if _blank(note):
        commercial = COMMERCIAL_APPEND
    elif COMMERCIAL_APPEND in str(note):
        commercial = note
    else:
        commercial = str(note).rstrip() + " " + COMMERCIAL_APPEND
    assumes = R.get(PID, "first_export_year_assumes_fid")
    if not _blank(assumes) and assumes is not False:
        fail(
            "first_export_year_assumes_fid is "
            f"{assumes!r}; this update leaves a blank or False value in place"
        )

    R.set(
        PID, "fid_confirmed", True,
        source=(
            "Final investment decision 29 Sep 2026. " + LNGC + " | " + SHELL
        ),
    )
    R.set(
        PID, "fid_date", date(2026, 9, 29),
        source=(
            "Announced 28 September 2026 Pacific time, 29 September 2026 UK time. "
            + LNGC + " | " + SHELL
        ),
    )
    R.set(
        PID, "tier", "under_construction",
        source="Moved from advanced_proposed after the 29 Sep 2026 final investment decision. " + LNGC,
    )
    R.set(
        PID, "calc_group", "under_construction",
        source="Follows tier after the 29 Sep 2026 final investment decision. Committed is operating plus under_construction.",
    )
    R.set(
        PID, "status", "construction",
        source=(
            "Post-FID value used by Cedar LNG and Woodfibre LNG. "
            "Field Definitions: GEM status vocabulary operating, construction, proposed, shelved, cancelled. "
            + LNGC
        ),
    )
    R.set(PID, "status_verified_date", "2026-09-29")
    R.set(PID, "tier_reason", TIER_REASON, source=TIER_REASON + " " + " | ".join(NEW_SOURCES))
    R.set(PID, "first_export_year", 2031, source=FIRST_EXPORT_SOURCE)
    R.set(
        PID, "commercial_note", commercial,
        source=COMMERCIAL_APPEND + " " + LNGC + " | " + SHELL + " | " + FT,
    )
    R.add_sources(PID, *NEW_SOURCES)


def update_coastal_gaslink(R: Register) -> None:
    ws = R.wb["Supporting Infrastructure"]
    headers = {c.value: c.column for c in ws[1] if c.value}
    target = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, headers["pipeline_id"]).value == "coastal_gaslink_expansion":
            target = r
            break
    if target is None:
        fail("Supporting Infrastructure has no coastal_gaslink_expansion row")
    notes_col = headers["notes"]
    src_col = headers["sources"]
    old_notes = ws.cell(target, notes_col).value or ""
    if CGL_NOTE not in str(old_notes):
        new_notes = (str(old_notes).rstrip() + " " + CGL_NOTE).strip()
        ws.cell(target, notes_col).value = new_notes
        R.log.append(("Supporting Infrastructure coastal_gaslink_expansion", "notes", old_notes, new_notes))
    old_src = ws.cell(target, src_col).value or ""
    if TC.split(": ")[-1] not in str(old_src):
        new_src = (str(old_src).rstrip() + " | " + TC).strip(" |")
        ws.cell(target, src_col).value = new_src
        R.log.append(("Supporting Infrastructure coastal_gaslink_expansion", "sources", old_src, new_src))


def _replace_text(wb, sheet: str, coordinate: str, old: str, new: str, log: list) -> None:
    cell = wb[sheet][coordinate]
    if cell.value != old:
        fail(f"{sheet}!{coordinate} is not the expected text:\n{cell.value!r}")
    cell.value = new
    log.append((f"{sheet}!{coordinate}", coordinate, old, new))


def update_wording(R: Register, data_inputs: Path):
    readme = R.wb["README"]
    b52 = readme["B52"].value
    old_b52_tail = (
        "Phase 2's 14 mtpa, 31 per cent of live export capacity, depends on the "
        "proposed second line agreed only in principle in March 2026."
    )
    new_b52_tail = (
        "Phase 2's 14 mtpa depends on the Coastal GasLink expansion, which will "
        "proceed following the LNG Canada Phase 2 final investment decision of "
        "29 September 2026."
    )
    if not isinstance(b52, str) or old_b52_tail not in b52:
        fail(f"README!B52 does not contain the pre-FID pipeline sentence:\n{b52!r}")
    new_b52 = b52.replace(old_b52_tail, new_b52_tail)
    readme["B52"].value = new_b52
    R.log.append(("Asset Register README!B52", "B52", b52, new_b52))

    _replace_text(
        R.wb, "README", "B60",
        "No FID yet. Status text refreshed; first LNG 2030 reconfirmed "
        "(GEM wiki, 25 Sep 2026: https://www.gem.wiki/LNG_Canada_Terminal).",
        "Final investment decision 29 Sep 2026. Status construction, tier and "
        "calc_group under_construction, first LNG 2031 (decision date plus the "
        "IEA 4 to 5 year lag; Shell expects commercial operations in the early 2030s).",
        R.log,
    )

    di = openpyxl.load_workbook(data_inputs)
    chains = di["Chains"]
    found = False
    for row in chains.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            if "65.7 proposed" not in cell.value:
                continue
            old = cell.value
            new = old.replace(
                "85.55 mtpa (14.0 operating, 5.85 under construction, 65.7 proposed; "
                "asset register of 27 Sep 2026)",
                "85.55 mtpa (14.0 operating, 19.85 under construction, 51.7 proposed; "
                "asset register of 29 Sep 2026)",
            )
            if new == old:
                fail(f"Chains!{cell.coordinate} mentions 65.7 proposed but not the expected sentence")
            cell.value = new
            R.log.append((f"Data Inputs Chains!{cell.coordinate}", cell.coordinate, old, new))
            found = True
    if not found:
        fail("Data Inputs Chains has no 65.7 proposed capacity note")
    return di


def assert_export_split(R: Register) -> None:
    totals = {k: 0.0 for k in EXPECTED_EXPORT}
    for pid, r in R.rows.items():
        chain = R.get(pid, "chain")
        group = R.get(pid, "calc_group")
        cap = R.get(pid, "capacity_mtpa")
        if chain != "export" or group not in totals or cap is None:
            continue
        totals[group] += float(cap)
    for group, expected in EXPECTED_EXPORT.items():
        if abs(totals[group] - expected) > 1e-6:
            fail(f"export {group} capacity is {totals[group]}, expected {expected}")
    total = sum(totals.values())
    if abs(total - 85.55) > 1e-6:
        fail(f"export capacity total is {total}, expected 85.55")


def main() -> None:
    if len(sys.argv) != 2:
        fail("usage: register_update_2026-09-29.py path/to/Canada_LNG_Asset_Register.xlsx")
    path = Path(sys.argv[1])
    data_inputs = path.parent / "Canada_LNG_Data_Inputs.xlsx"
    if not path.is_file():
        fail(f"register not found: {path}")
    if not data_inputs.is_file():
        fail(f"data inputs not found: {data_inputs}")
    R = Register(str(path))
    assert_current(R)
    update_phase_2(R)
    update_coastal_gaslink(R)
    data_wb = update_wording(R, data_inputs)
    assert_export_split(R)
    # GEM's last check predates the decision. Leave both cells alone.
    if R.get(PID, "gem_status_verbatim") != "proposed":
        fail("gem_status_verbatim changed")
    if str(R.get(PID, "gem_status_date"))[:10] != "2026-08-28":
        fail(f"gem_status_date changed: {R.get(PID, 'gem_status_date')!r}")
    R.wb.save(path)
    data_wb.save(data_inputs)
    data_wb.close()
    print(f"updated {path}")
    print(f"updated {data_inputs}")
    for where, field, old, new in R.log:
        print(f"CHANGED {where} [{field}]")
        print(f"  before: {old!r}")
        print(f"  after:  {new!r}")


if __name__ == "__main__":
    main()
