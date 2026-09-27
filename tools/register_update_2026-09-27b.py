"""Asset register update 27 Sep 2026, part b: first_export_year_assumes_fid.

Adds the register field that tells the model not to add the five-year FID delay
(fid_delay_mid) to a pre-FID asset whose first_export_year is a developer date
that already allows for the build after an FID. Decision: Freddie, 27 Sep 2026
("skip those delays in those cases"). Code side: src/model.py _date_assumes_fid /
_fid_delay_applies, used by src/trajectories.py and src/monte_carlo.py.

Usage: python register_update_2026-09-27b.py path/to/Canada_LNG_Asset_Register.xlsx
"""

from __future__ import annotations

import sys
from copy import copy

import openpyxl

MPO_KSI = "https://www.canada.ca/en/privy-council/major-projects-office/projects/national/ksi-lisims.html"
GEM_KSI = "https://www.gem.wiki/Ksi_Lisims_FLNG_Terminal"
BC_T1B = "https://news.gov.bc.ca/releases/2026ECS0043-000870"

FIELD = "first_export_year_assumes_fid"
FLAGS = {
    "ksi_lisims_lng": (
        f"Major Projects Office, Ksi Lisims LNG (modified 14 Sep 2026): {MPO_KSI} - Uniper "
        "agreement of 29 Jul 2026, 'first deliveries expected to begin 2032'. GEM wiki: "
        f"{GEM_KSI} - FID 'expected in 2026'. The 2032 date is a developer delivery date that "
        "already allows for the build after a 2026 FID, so fid_delay_mid is not added on top "
        "(decision 27 Sep 2026)."
    ),
    "tilbury_phase_1b": (
        f"BC Government news release 2026ECS0043 (24 Jul 2026): {BC_T1B} - construction 'as "
        "early as mid-2027', 'in service as early as 2031'. The 2031 date already includes the "
        "build period, so fid_delay_mid is not added on top (decision 27 Sep 2026)."
    ),
}
NOT_FLAGGED = ("Not applicable. Blank = False: the asset has an FID, or its first_export_year "
               "is not a developer date that already allows for the build after an FID, so "
               "fid_delay_mid applies as before.")


def add_column(ws, header: str) -> int:
    last = ws.max_column
    col = last + 1
    src_h, dst_h = ws.cell(1, last), ws.cell(1, col)
    dst_h.value = header
    for attr in ("font", "fill", "border", "alignment", "number_format", "protection"):
        setattr(dst_h, attr, copy(getattr(src_h, attr)))
    letter_src = openpyxl.utils.get_column_letter(last)
    letter_dst = openpyxl.utils.get_column_letter(col)
    ws.column_dimensions[letter_dst].width = ws.column_dimensions[letter_src].width
    for r in range(2, ws.max_row + 1):
        s, d = ws.cell(r, last), ws.cell(r, col)
        for attr in ("font", "fill", "border", "alignment", "protection"):
            setattr(d, attr, copy(getattr(s, attr)))
    return col


def main(path: str):
    wb = openpyxl.load_workbook(path)
    reg, src = wb["Asset Register"], wb["Asset Sources"]
    assert FIELD not in [c.value for c in reg[1]], "field already present"
    cr, cs = add_column(reg, FIELD), add_column(src, FIELD)
    for r in range(2, reg.max_row + 1):
        pid = reg.cell(r, 1).value
        if not pid:
            continue
        assert src.cell(r, 1).value == pid
        if pid in FLAGS:
            reg.cell(r, cr).value = True
            src.cell(r, cs).value = FLAGS[pid]
            print(f"{pid}: {FIELD} = True")
        else:
            src.cell(r, cs).value = NOT_FLAGGED

    fd = wb["Field Definitions"]
    last = fd.max_row
    new = last + 1
    for c in range(1, fd.max_column + 1):
        s, d = fd.cell(last, c), fd.cell(new, c)
        for attr in ("font", "fill", "border", "alignment", "protection"):
            setattr(d, attr, copy(getattr(s, attr)))
    vals = [FIELD, "Asset Register",
            "True when first_export_year is a developer date that already allows for the "
            "build after an FID (for example a delivery start in an offtake agreement). The "
            "model then does not add fid_delay_mid to that pre-FID asset. Blank means False.",
            "-", "assigned",
            "Judgement from the source of first_export_year, stated in the Asset Sources cell. "
            "Added 27 Sep 2026."]
    for c, v in enumerate(vals, start=1):
        fd.cell(new, c).value = v

    rd = wb["README"]
    r = rd.max_row + 1
    for c in (1, 2):
        s, d = rd.cell(r - 1, c), rd.cell(r, c)
        for attr in ("font", "fill", "border", "alignment"):
            setattr(d, attr, copy(getattr(s, attr)))
    rd.cell(r, 1).value = "FID delay flag"
    rd.cell(r, 2).value = (
        f"New field {FIELD}. Where True, the model does not add the five-year FID delay to a "
        "pre-FID asset because its first_export_year already allows for the build: Ksi Lisims "
        f"(2032, Uniper agreement: {MPO_KSI}) and Tilbury Phase 1b (2031, {BC_T1B}).")
    wb.save(path)


if __name__ == "__main__":
    main(sys.argv[1])
