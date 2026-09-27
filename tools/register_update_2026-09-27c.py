"""Asset register update 27 Sep 2026, part c: new rows with verified sources.

Adds four rows (items 16, 17, 18 and 23 of the model update list). None enters
the headline: three are watch rows (excluded from all totals) and one is a
domestic plant with no published capacity. Dawson Creek, Elmworth, Hagar and the
Linde demonstration are held: their sources could not be opened or did not
confirm operation (Campus Energy's site no longer lists Dawson Creek).

Usage: python register_update_2026-09-27c.py path/to/Canada_LNG_Asset_Register.xlsx
"""

from __future__ import annotations

import sys
from copy import copy

import openpyxl

NP_WEB = "https://www.northernprincelng.com/about"
NEESTANAN = "https://neestanan.ca/neestanan-and-lng-exports/"
BOE_2024 = "https://boereport.com/2024/03/05/neestanan-signs-mou-with-northern-prince-lng/"
BOE_2025 = "https://boereport.com/2025/07/02/neestanan-and-fox-lake-cree-nation-secure-federal-authorization-to-export-lng/"
REGDOCS_NP = "https://apps.cer-rec.gc.ca/REGDOCS/Item/View/4459935"
MPO_CHURCHILL = "https://www.canada.ca/en/privy-council/major-projects-office/projects/other/referred/churchill-plus.html"
PROSPECTUS = "https://cabinradio.ca/wp-content/uploads/2026/09/CanadaInvestmentSummit.pdf"
CBC_CHURCHILL = "https://ca.news.yahoo.com/port-churchill-railway-improvements-cost-100000703.html"
DEEPDIVE = "https://thedeepdive.ca/manitoba-pitches-pst-exemption-to-drive-port-of-churchill-investment/"
GLOBAL_NEWS = "https://globalnews.ca/news/11806964/federal-government-puts-timeline-in-place-for-port-of-churchill-project-kinew/"
CBC_NEESTANAN = "https://www.cbc.ca/1.7582120"

NA, NAPP = "Not available", "Not applicable"
WATCH_BASIN = "not available - watch row, excluded from all totals; no feedgas route stated"
NOT_FLAGGED = ("Not applicable. Blank = False: the asset has an FID, or its first_export_year "
               "is not a developer date that already allows for the build after an FID, so "
               "fid_delay_mid applies as before.")

ROWS = {
    "northern_prince_lng": {
        "project_name": ("Northern Prince LNG", "Northern Prince LNG"),
        "tier": ("watch", "Assigned by us - see tier_reason"),
        "calc_group": ("watch", "Assigned: watch rows are excluded from all totals"),
        "tier_reason": (
            "Not a live export proposal: no filed project description or environmental "
            "assessment, no named buyer and no named site. The developer's website claims heads "
            "of agreement for 6 mtpa on 'our larger project' and an 'initial LNG Export Permit'; "
            "the only CER records located are applications for short-term natural gas export "
            "orders and Order GO-116-2022, which authorise no facility or long-term volume. "
            "NeeStaNan places the 6 mtpa facility near Summit Lake, BC, the same area and "
            "containerised-rail concept as Summit Lake PG LNG, so the two may overlap. Watch "
            "rows enter no total.",
            f"Northern Prince LNG website: {NP_WEB} | NeeStaNan: {NEESTANAN} | BOE Report, "
            f"5 Mar 2024: {BOE_2024} | CER REGDOCS filing titles (documents not opened): {REGDOCS_NP}"),
        "province": ("British Columbia", f"Northern Prince LNG website: {NP_WEB}"),
        "location": ("Near Summit Lake, north of Prince George (per NeeStaNan); the developer names no site",
                     f"NeeStaNan: {NEESTANAN} - 'near Summit Lake, British Columbia'"),
        "proponent": ("Northern Prince LNG Inc.", f"Northern Prince LNG website: {NP_WEB}"),
        "owners": ("Not disclosed ('Canadian Federally incorporated Company')", f"Northern Prince LNG website: {NP_WEB}"),
        "status": ("proposed", f"Northern Prince LNG website: {NP_WEB}"),
        "status_verified_date": ("2026-09", "Date we last verified against the sources listed"),
        "fid_confirmed": (0, NA),
        "export_licence": ("Short-term export order GO-116-2022 (term and volume not verified); no long-term GL licence located",
                           f"CER REGDOCS: {REGDOCS_NP} - filing titles 'Application for Short-Term Natural Gas "
                           "Export Orders' (C20575, C29880) and 'Order GO-116-2022'; documents not opened (REGDOCS unreachable)"),
        "capacity_mtpa": (6.0, f"Northern Prince LNG website: {NP_WEB} - 'executed HOA's for 6MMTPA for our larger project'; "
                               f"NeeStaNan: {NEESTANAN} - '6 million tonne per annum (mtpa) LNG manufacturing facility'"),
        "capacity_basis": ("Developer claim: heads of agreement for 6 mtpa on 'our larger project'; counterparties not "
                           "named; no filed figure. Watch row, so it enters no total.", f"Northern Prince LNG website: {NP_WEB}"),
        "end_use": ("overseas export", "Derived: classified as an export project"),
        "domestic_supply_mtpa": (0, "Derived: export projects supply no LNG to the domestic market"),
        "domestic_supply_note": ("Export project.", "Compiled by us"),
        "destination_market": ("Asia and Europe", f"Northern Prince LNG website: {NP_WEB}"),
        "export_route": ("Containerised LNG by rail or truck to the coast (BOE Report, Mar 2024); port not named",
                         f"BOE Report, 5 Mar 2024: {BOE_2024}"),
        "commercial_note": (
            "MOU with NeeStaNan (Fox Lake Cree Nation) to assess a mid-scale LNG production and "
            "export operation at a Hudson Bay port; BOE Report dates it March 2024, a July 2025 "
            "release gives 8 Feb 2025. The website also mentions 'the first of a number of "
            "small-scale projects'.",
            f"BOE Report, 5 Mar 2024: {BOE_2024} | BOE Report, 2 Jul 2025: {BOE_2025} | Website: {NP_WEB}"),
        "facility_type": ("export (not specified)", f"Northern Prince LNG website: {NP_WEB}"),
        "liquefaction_drive": ("not_published", "Not published"),
        "liquefaction_drive_note": ("Drive type not published.", "Compiled by us"),
        "chain": ("export", "Assigned by us: stated export project. Watch rows are excluded from all totals."),
        "chain_note": ("Assigned from end_use. Watch row: excluded from all totals.", "Compiled by us"),
        "export_term_note": ("Short-term order only; no long-term licence.", f"CER REGDOCS: {REGDOCS_NP}"),
        "gem_project_id": (None, "Not available - does not appear in GEM GGIT LNG Terminals Sep 2025"),
        "nrcan_listed": ("no", "NRCan Canadian LNG projects: not present (page modified 7 Jan 2025)"),
        "cer_export_licence": ("GO-116-2022 (short-term order)", f"CER REGDOCS: {REGDOCS_NP}"),
        "cer_licence_url": (REGDOCS_NP, "CER REGDOCS"),
        "sources": (f"Northern Prince LNG website: {NP_WEB} | NeeStaNan, NeeStaNan and LNG Exports: {NEESTANAN} | "
                    f"BOE Report, NeeStaNan signs MOU with Northern Prince LNG (5 Mar 2024): {BOE_2024} | "
                    f"BOE Report, NeeStaNan and Fox Lake Cree Nation secure federal authorization (2 Jul 2025): {BOE_2025} | "
                    f"CER REGDOCS, Northern Prince LNG Inc.: {REGDOCS_NP}",
                    "Compiled by us from the sources named in this row"),
        "feedgas_basin": ("not available", WATCH_BASIN),
        "gem_status_verbatim": (None, "not available - no GEM entry"),
        "gem_status_date": (None, "not available - no GEM entry"),
    },
    "hudson_bay_churchill_lng": {
        "project_name": ("Port of Churchill Plus LNG (floating LNG concept)", "Port of Churchill Plus LNG (floating LNG concept)"),
        "tier": ("watch", "Assigned by us - see tier_reason"),
        "calc_group": ("watch", "Assigned: watch rows are excluded from all totals"),
        "tier_reason": (
            "Government concept with no developer and no capacity. Port of Churchill Plus was "
            "referred to the Major Projects Office on 11 Sep 2025 as a transformative strategy; "
            "the MPO page (modified 21 May 2026) does not mention LNG. The federal Canada "
            "Investment Summit Prospectus (Sep 2026, p. 40) lists an 'LNG export corridor' with "
            "US$57B capex at the planning, engineering and feasibility stage. Manitoba's premier "
            "costs a version with a floating offshore LNG platform at $70 to 80 billion (CBC, "
            "15 Sep 2026); Manitoba's sales tax exemption announced 14 Sep 2026 covers "
            "'liquefied natural gas facilities'; Ottawa wants LNG shipped by 2030. The port is "
            "ice-bound outside late July to early November.",
            f"Major Projects Office, Port of Churchill Plus: {MPO_CHURCHILL} | Canada Investment Summit "
            f"Prospectus, p. 40: {PROSPECTUS} | CBC via Yahoo, 15 Sep 2026: {CBC_CHURCHILL} | The Deep Dive, "
            f"15 Sep 2026: {DEEPDIVE} | Global News, 17 Apr 2026: {GLOBAL_NEWS}"),
        "province": ("Manitoba", f"Major Projects Office: {MPO_CHURCHILL}"),
        "location": ("Port of Churchill, Hudson Bay (offshore floating platform concept)", f"CBC via Yahoo, 15 Sep 2026: {CBC_CHURCHILL}"),
        "proponent": ("None named; promoted by the Government of Manitoba", f"CBC via Yahoo, 15 Sep 2026: {CBC_CHURCHILL}; The Deep Dive: {DEEPDIVE}"),
        "owners": (None, NA),
        "status": ("proposed", f"Canada Investment Summit Prospectus, p. 40: {PROSPECTUS}"),
        "status_verified_date": ("2026-09", "Date we last verified against the sources listed"),
        "fid_confirmed": (0, NA),
        "capacity_mtpa": (None, "Not available - no source states a capacity; blank rather than estimated"),
        "capacity_basis": ("Not available", "Not available - no source states a capacity"),
        "end_use": ("overseas export", "Derived: classified as an export concept"),
        "destination_market": ("Not stated", NA),
        "commercial_note": ("Federal target: LNG shipments by 2030, or federal backing is at risk (Global News, 17 Apr 2026).",
                            f"Global News, 17 Apr 2026: {GLOBAL_NEWS}"),
        "facility_type": ("export, floating (concept)", f"CBC via Yahoo, 15 Sep 2026: {CBC_CHURCHILL}"),
        "floating": ("yes", f"CBC via Yahoo, 15 Sep 2026: {CBC_CHURCHILL}"),
        "liquefaction_drive": ("not_published", "Not published"),
        "liquefaction_drive_note": ("Drive type not published.", "Compiled by us"),
        "chain": ("export", "Assigned by us: export concept. Watch rows are excluded from all totals."),
        "chain_note": ("Assigned from end_use. Watch row: excluded from all totals.", "Compiled by us"),
        "export_term_note": ("No export licence identified.", NA),
        "gem_project_id": (None, "Not available - does not appear in GEM GGIT LNG Terminals Sep 2025"),
        "nrcan_listed": ("no", "NRCan Canadian LNG projects: not present (page modified 7 Jan 2025)"),
        "regulatory_profile_url": (MPO_CHURCHILL, "Major Projects Office"),
        "sources": (f"Major Projects Office, Port of Churchill Plus: {MPO_CHURCHILL} | Invest in Canada, Canada Investment "
                    f"Summit Prospectus (Sep 2026), p. 40: {PROSPECTUS} | CBC via Yahoo, Port of Churchill cost (15 Sep 2026): "
                    f"{CBC_CHURCHILL} | The Deep Dive, Manitoba PST exemption (15 Sep 2026): {DEEPDIVE} | Global News, "
                    f"federal timeline for Port of Churchill (17 Apr 2026): {GLOBAL_NEWS}",
                    "Compiled by us from the sources named in this row"),
        "feedgas_basin": ("not available", WATCH_BASIN),
        "gem_status_verbatim": (None, "not available - no GEM entry"),
        "gem_status_date": (None, "not available - no GEM entry"),
    },
    "neestanan_port_nelson_lng": {
        "project_name": ("NeeStaNan Port Nelson LNG export (feasibility)", "NeeStaNan Port Nelson LNG export (feasibility)"),
        "tier": ("watch", "Assigned by us - see tier_reason"),
        "calc_group": ("watch", "Assigned: watch rows are excluded from all totals"),
        "tier_reason": (
            "Feasibility study for a multi-commodity port at Port Nelson that would include LNG. "
            "The CER authorised NeeStaNan to export LNG until June 2027 so it can carry out the "
            "study (announced 2 Jul 2025). No capacity, facility filing or environmental "
            "assessment. Manitoba's premier: 'We haven't actually done an announcement of this "
            "thing yet.' MOU with Northern Prince LNG.",
            f"CBC: {CBC_NEESTANAN} | BOE Report, 2 Jul 2025: {BOE_2025}"),
        "province": ("Manitoba", f"CBC: {CBC_NEESTANAN}"),
        "location": ("Proposed Port Nelson at the mouth of the Nelson River, Hudson Bay; planned 150 km rail spur from Gillam",
                     f"CBC: {CBC_NEESTANAN}"),
        "proponent": ("NeeStaNan", f"CBC: {CBC_NEESTANAN}"),
        "owners": ("NeeStaNan; majority owner Fox Lake Cree Nation", f"BOE Report, 2 Jul 2025: {BOE_2025}"),
        "status": ("proposed", f"CBC: {CBC_NEESTANAN} - feasibility study phase"),
        "status_verified_date": ("2026-09", "Date we last verified against the sources listed"),
        "fid_confirmed": (0, NA),
        "export_licence": ("Short-term CER export authorisation to June 2027, for a feasibility study (order number and volume not verified)",
                           f"CBC: {CBC_NEESTANAN}; BOE Report, 2 Jul 2025: {BOE_2025}"),
        "capacity_mtpa": (None, "Not available - no capacity stated; blank rather than estimated"),
        "capacity_basis": ("Not available", "Not available - no capacity stated"),
        "end_use": ("overseas export", "Derived: classified as an export concept"),
        "destination_market": ("Not stated", NA),
        "commercial_note": ("MOU with Northern Prince LNG to assess a mid-scale LNG production and export operation (BOE Report).",
                            f"BOE Report, 2 Jul 2025: {BOE_2025}"),
        "facility_type": ("multi-commodity port including LNG (concept)", f"CBC: {CBC_NEESTANAN}"),
        "liquefaction_drive": ("not_published", "Not published"),
        "liquefaction_drive_note": ("Drive type not published.", "Compiled by us"),
        "chain": ("export", "Assigned by us: export concept. Watch rows are excluded from all totals."),
        "chain_note": ("Assigned from end_use. Watch row: excluded from all totals.", "Compiled by us"),
        "export_term_note": ("Short-term authorisation to June 2027 only.", f"CBC: {CBC_NEESTANAN}"),
        "gem_project_id": (None, "Not available - does not appear in GEM GGIT LNG Terminals Sep 2025"),
        "nrcan_listed": ("no", "NRCan Canadian LNG projects: not present (page modified 7 Jan 2025)"),
        "sources": (f"CBC, CER authorises NeeStaNan LNG exports: {CBC_NEESTANAN} | BOE Report, NeeStaNan and Fox Lake Cree "
                    f"Nation secure federal authorization (2 Jul 2025): {BOE_2025} | NeeStaNan, NeeStaNan and LNG Exports: {NEESTANAN}",
                    "Compiled by us from the sources named in this row"),
        "feedgas_basin": ("not available", WATCH_BASIN),
        "gem_status_verbatim": (None, "not available - no GEM entry"),
        "gem_status_date": (None, "not available - no GEM entry"),
    },
    "project_northern_lights_lng": {
        "project_name": ("Project Northern Lights (Cool LNG)", "Project Northern Lights (Cool LNG)"),
        "tier": ("advanced_proposed", "Assigned by us - see tier_reason"),
        "calc_group": ("proposed", "Assigned: proposed"),
        "tier_reason": (
            "Shovel-ready small-scale LNG plant, the first commercial deployment of Cool LNG's MA3 "
            "liquefaction technology: land and gas supply secured, engineering, zoning and "
            "subdivision completed and 'all required permits obtained'. Stage 1 capex US$60M, "
            "US$250M for stages 1 to 5; seeking equity, strategic partners and long-term offtake. "
            "No capacity or site published, so it enters no total.",
            f"Canada Investment Summit Prospectus (Sep 2026), p. 7: {PROSPECTUS}"),
        "province": ("Alberta", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "location": ("Alberta (site not published)", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "proponent": ("Cool LNG (contact at Cool Ventures Inc.)", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "owners": (None, NA),
        "status": ("proposed", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS} - 'Shovel-ready'"),
        "status_verified_date": ("2026-09", "Date we last verified against the sources listed"),
        "fid_confirmed": (0, "Not available - no FID stated"),
        "export_licence": ("Not applicable - domestic facility", NAPP),
        "capacity_mtpa": (None, "Not available - capacity not published; blank rather than estimated"),
        "capacity_basis": ("Not available - the prospectus gives cost but no capacity", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "end_use": ("domestic supply (industrial, mining, remote communities, transport, utilities)",
                    f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "end_use_note": ("Markets across Western Canada.", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "domestic_supply_note": ("Small-scale domestic supply; capacity not published.", "Compiled by us"),
        "destination_market": ("domestic", "Derived: domestic facility"),
        "export_route": ("Not applicable", NAPP),
        "offtake_contracted_mtpa": (None, NA),
        "facility_type": ("small-scale liquefaction", f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "liquefaction_drive": ("not_published", "Not published for small plants"),
        "liquefaction_drive_note": ("Drive type not published. Uses Cool LNG's proprietary MA3 liquefaction technology.",
                                    f"Canada Investment Summit Prospectus, p. 7: {PROSPECTUS}"),
        "chain": ("domestic", "Assigned by us from end_use"),
        "chain_note": ("Assigned from end_use. Outside the headline export scope.", "Compiled by us"),
        "authorised_export_term_years": (None, NAPP),
        "export_term_note": ("No export licence - domestic facility", NAPP),
        "gem_project_id": (None, "Not available - does not appear in GEM GGIT LNG Terminals Sep 2025"),
        "nrcan_listed": ("no", "NRCan Canadian LNG projects: not present (page modified 7 Jan 2025)"),
        "sources": (f"Invest in Canada, Canada Investment Summit Prospectus (Sep 2026), p. 7: {PROSPECTUS}",
                    "Compiled by us from the sources named in this row"),
        "feedgas_basin": (None, "not available - no capacity, so no lifecycle is computed"),
        "gem_status_verbatim": (None, "not available - no GEM entry"),
        "gem_status_date": (None, "not available - no GEM entry"),
    },
}

GAPS = [
    ("verification", "Northern Prince LNG", "Watch; 6 mtpa developer claim",
     "No filed project description, named buyer, named site or long-term CER licence. The only CER "
     "records are short-term export order filings (titles only). Possible overlap with Summit Lake PG LNG.",
     "A filed project description, a named HOA buyer, a named site or a long-term licence application",
     "medium - as a proposed row 6 mtpa would add about 610 MtCO2e"),
    ("capacity_mtpa / proponent", "Port of Churchill Plus LNG", "Blank",
     "No developer and no stated capacity; the LNG element is a government concept.",
     "A named developer and a stated liquefaction capacity", "medium - federal 2030 target"),
    ("capacity_mtpa", "NeeStaNan Port Nelson", "Blank",
     "Feasibility study under a short-term CER authorisation; no capacity stated.",
     "Feasibility study results or a filed project description", "low"),
    ("capacity_mtpa / location", "Project Northern Lights (Cool LNG)", "Blank",
     "The federal prospectus gives cost and permitting status but no capacity or site.",
     "Developer statement of capacity and site, or the Alberta permit record", "low - domestic, outside the headline"),
    ("verification", "Dawson Creek LNG, Elmworth LNG, Hagar LNG, Linde Fort Saskatchewan", "Not added",
     "Candidate domestic plants from the 26 Sep 2026 research sweep. Their sources could not be opened, "
     "or did not confirm current operation (Campus Energy's site no longer lists Dawson Creek).",
     "An operator page or regulator record confirming operation and capacity", "low - domestic, outside the headline"),
]


def copy_style(src_cell, dst_cell):
    for attr in ("font", "fill", "border", "alignment", "number_format", "protection"):
        setattr(dst_cell, attr, copy(getattr(src_cell, attr)))


def main(path: str):
    wb = openpyxl.load_workbook(path)
    reg, src = wb["Asset Register"], wb["Asset Sources"]
    cols = {c.value: c.column for c in reg[1] if c.value}
    scols = {c.value: c.column for c in src[1] if c.value}
    existing = {reg.cell(r, 1).value for r in range(2, reg.max_row + 1)}
    template = reg.max_row
    for pid, fields in ROWS.items():
        assert pid not in existing, pid
        unknown = set(fields) - set(cols)
        assert not unknown, unknown
        r = reg.max_row + 1
        for name, c in cols.items():
            copy_style(reg.cell(template, c), reg.cell(r, c))
            copy_style(src.cell(template, scols[name]), src.cell(r, scols[name]))
            if name == "project_id":
                val, source = pid, pid
            elif name in fields:
                val, source = fields[name]
            elif name == "first_export_year_assumes_fid":
                val, source = None, NOT_FLAGGED
            else:
                val, source = None, NA
            reg.cell(r, c).value = val
            src.cell(r, scols[name]).value = source
        print(f"added {pid} at row {r}")

    dg = wb["Data Gaps"]
    for g in GAPS:
        r = dg.max_row + 1
        for c, v in enumerate(g, start=1):
            copy_style(dg.cell(r - 1, c), dg.cell(r, c))
            dg.cell(r, c).value = v

    rd = wb["README"]
    for label, text in [
        ("New rows 27 Sep 2026",
         "Northern Prince LNG, Port of Churchill Plus LNG and NeeStaNan Port Nelson added as watch rows "
         "(excluded from all totals); Project Northern Lights added as a domestic row with no published "
         "capacity (enters no total). Dawson Creek, Elmworth, Hagar and Linde held until sources confirm them."),
    ]:
        r = rd.max_row + 1
        for c in (1, 2):
            copy_style(rd.cell(r - 1, c), rd.cell(r, c))
        rd.cell(r, 1).value, rd.cell(r, 2).value = label, text
    for r in range(1, rd.max_row + 1):
        if rd.cell(r, 1).value == "watch":
            rd.cell(r, 2).value = (rd.cell(r, 2).value.rstrip()
                                   + " Northern Prince LNG and the Churchill and Port Nelson Hudson Bay LNG concepts are also watch rows.")
    wb.save(path)


if __name__ == "__main__":
    main(sys.argv[1])
