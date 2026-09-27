"""Asset register update, 27 September 2026.

Applies the "Register: existing rows" changes (items 1-15 of the model update
list) to Inputs/Canada_LNG_Asset_Register.xlsx. Every changed value gets a
matching cell in the Asset Sources sheet with its source and link.

Usage: python register_update_2026-09-27.py path/to/Canada_LNG_Asset_Register.xlsx

Decisions behind it (Freddie, 27 Sep 2026): do not wait for the LNG Canada
Phase 2 FID; take the developer's figure for Fermeuse; use the latest
permitted figure for Cedar; add the new Kino Aski sources; and, generally, use
the latest available data, sourced and linked, for anything that moves figures.
Kitsault is not changed: both BusinessWire sources returned 404 when checked.
"""

from __future__ import annotations

import sys
from copy import copy

import openpyxl

# ---------------------------------------------------------------- sources
MPO_P2 = "https://www.canada.ca/en/privy-council/major-projects-office/projects/national/lng-canada.html"
GEM_LNGC = "https://www.gem.wiki/LNG_Canada_Terminal"
NV = "https://thenorthernview.com/2026/09/24/lng-canada-phase-2-decision-could-come-as-early-as-october/"
MPO_KSI = "https://www.canada.ca/en/privy-council/major-projects-office/projects/national/ksi-lisims.html"
BC_SEFE = "https://news.gov.bc.ca/releases/2026ECS0025-000604"
NRCAN_SANTOS = "https://www.canada.ca/en/natural-resources-canada/news/2026/09/canadian-lng-project-signs-third-international-offtake-deal.html"
GEM_KSI = "https://www.gem.wiki/Ksi_Lisims_FLNG_Terminal"
BC_NCTL = "https://news.gov.bc.ca/releases/2026PREM0034-001026"
BC_HYDRO_KSI = "https://news.gov.bc.ca/releases/2026ECS0002-000043"
NGI_KSI = "https://naturalgasintel.com/news/ksi-lisims-lng-power-agreement-advances-bcs-natural-gas-export-buildout/"
TOTAL_KSI = "https://www.businesswire.com/news/home/20250519012495/en/"
NARWHAL_F = "https://thenarwhal.ca/lng-newfoundland-lessons-kitimat-bc/"
GEM_FERM = "https://www.gem.wiki/Fermeuse_Energy_FLNG_Terminal"
EAO_CEDAR = ("https://projects.eao.gov.bc.ca/api/public/document/68caf0067a03b4002205daf9/"
             "download/Cedar%20LNG%20Accommodation%20+%20Capacity%20Amendment%20Application.pdf")
GIIGNL = "https://www.giignl.org/news/cedar-lng-receives-approval-to-increase-project-capacity"
VANGUARD = "https://www.thecanadianvanguard.com/cedar-lng-secures-bc-regulatory-nod-to-expand-future-output-capacity/"
CEDAR_FID = "https://www.cedarlng.com/cedar-lng-announces-positive-final-investment-decision/"
PEMBINA_DEC25 = "https://www.pembina.com/media-centre/news/details/26400a51-ee23-4209-b789-4a15c2f699d1"
CEDAR_SPRING26 = "https://www.cedarlng.com/2026/04/30/spring-2026-newsletter/"
KINO_WEB = "https://kinoaskilng.ca/"
KINO_NEWSWIRE = "https://www.newswire.ca/news-releases/first-nations-lead-kino-aski-lng-a-major-canadian-energy-project-800865754.html"
GASWORLD = "https://www.gasworld.com/story/kino-aski-lng-targets-north-america-europe-energy-corridor/2257622.article/"
PROSPECTUS = "https://cabinradio.ca/wp-content/uploads/2026/09/CanadaInvestmentSummit.pdf"
ENERGYMIX_KINO = "https://www.theenergymix.com/kino-aski-lng-signs-tentative-deal-with-ukraine-state-fossil-as-opposition-rises-complications-abound/"
QN = "https://www.quebecnouvelles.com/un-troisieme-trace-etudie-pour-livrer-du-gnl-a-baie-comeau-1441870.html"
HANWHA_KANATA = "https://www.newswire.ca/news-releases/hanwha-ocean-signs-strategic-memorandum-of-understanding-with-kanata-clean-power-for-proposed-floating-lng-project-in-canada-867061587.html"
LNGI_KANATA = "https://www.lngindustry.com/floating-lng/14092026/kanata-clean-power-and-hanwha-ocean-to-launch-pre-feed-for-proposed-floating-lng-project/"
LNGJ_WF = "https://lngjournal.com/index.php/latest-news-mainmenu-47/item/116798-woodfibre-lng-on-track-to-complete-construction-by-late-2027"
WF_CONSTR = "https://woodfibrelng.ca/construction/"
ENERGETIC = "https://energeticcity.ca/2025/08/29/export-licence-extension-for-woodfibre-lng-granted-by-energy-regulator/"
GEM_GGIT = "https://globalenergymonitor.org/projects/global-gas-infrastructure-tracker/"
IAAC_TIL = "https://www.canada.ca/en/impact-assessment-agency/news/2026/09/major-milestone-completed.html"
SURREY = "https://surreynowleader.com/2026/09/21/delta-lng-expansion-project-issued-b-c-environmental-assessment-certificate/"
BC_T1B = "https://news.gov.bc.ca/releases/2026ECS0043-000870"
NRCAN_LIST = "https://natural-resources.canada.ca/energy-sources/fossil-fuels/canadian-liquified-natural-gas-projects"
TILPAC = "https://tilburypacific.ca/about-tilbury-pacific/"
YAQWA = "https://www.yaqwadevcorp.com/2025/11/18/yaq%CA%B7a-acquisition/"
EPIC_KIT = "https://projects.eao.gov.bc.ca/api/public/search?dataset=Project&keywords=Kitimat&pageNum=0&pageSize=100"
GEM_KIT = "https://www.gem.wiki/Kitimat_LNG_Terminal"
NL_EA = "https://www.gov.nl.ca/ecc/env-assessment/projects-list/"
CRYOPEAK = "https://www.cryopeak.com/"

VERIFIED = "2026-09"


class Register:
    def __init__(self, path: str):
        self.path = path
        self.wb = openpyxl.load_workbook(path)
        self.reg = self.wb["Asset Register"]
        self.src = self.wb["Asset Sources"]
        self.cols = {c.value: c.column for c in self.reg[1] if c.value}
        # Same field names on both sheets, but four columns sit in a different order.
        self.scols = {c.value: c.column for c in self.src[1] if c.value}
        assert set(self.cols) == set(self.scols), "Asset Register and Asset Sources fields differ"
        self.rows = {}
        for r in range(2, self.reg.max_row + 1):
            pid = self.reg.cell(r, 1).value
            if pid:
                self.rows[pid] = r
                assert self.src.cell(r, 1).value == pid, f"row mismatch for {pid}"
        self.log: list[tuple[str, str, object, object]] = []

    def set(self, pid: str, field: str, value, source: str | None = None):
        r, c = self.rows[pid], self.cols[field]
        old = self.reg.cell(r, c).value
        self.reg.cell(r, c).value = value
        if source is not None:
            self.src.cell(r, self.scols[field]).value = source
        self.log.append((pid, field, old, value))

    def get(self, pid: str, field: str):
        return self.reg.cell(self.rows[pid], self.cols[field]).value

    def add_sources(self, pid: str, *entries: str):
        cur = self.get(pid, "sources") or ""
        parts = [p.strip() for p in cur.split("|") if p.strip()]
        for e in entries:
            url = e.split(": ")[-1].strip()
            if not any(url in p for p in parts):
                parts.append(e)
        self.set(pid, "sources", " | ".join(parts))

    def verified(self, pid: str):
        self.set(pid, "status_verified_date", VERIFIED)


def update_asset_rows(R: Register):
    # 1. LNG Canada Phase 2: no FID yet; status text and sources refreshed.
    p = "lng_canada_phase_2"
    R.set(p, "tier_reason",
          "Referred to the Major Projects Office Sep 2025 and designated national interest. "
          "Canada, British Columbia and LNG Canada agreed enhanced investment co-operation on "
          "14 May 2026, targeting a 2026 FID. LNG Canada issued a limited notice to proceed to "
          "the JGC-Fluor joint venture in June 2026. Reuters reported on 24 Sep 2026 that an FID "
          "could come as early as October, and LNG Canada hopes to decide before the end of "
          "2026. No FID as of 27 Sep 2026, so the tier stays advanced_proposed.",
          f"Major Projects Office, LNG Canada Phase 2 (modified 2 Jul 2026): {MPO_P2} | "
          f"GEM wiki, LNG Canada Terminal (edited 25 Sep 2026): {GEM_LNGC} | "
          f"The Northern View, 24 Sep 2026 (Reuters): {NV}")
    R.set(p, "first_export_year", 2030,
          f"GEM wiki, LNG Canada Terminal (edited 25 Sep 2026): {GEM_LNGC} - planned first LNG "
          f"for Phase 2 is 2030. Unchanged from GEM GGIT (Sep 2025): {GEM_GGIT} "
          "LatestPlannedStartYear=2030. No developer date published; replace with the date "
          "given at FID.")
    R.add_sources(p, f"The Northern View, Phase 2 decision could come as early as October (24 Sep 2026): {NV}")
    R.verified(p)

    # 2. Ksi Lisims: first deliveries 2032; offtake; licence clause; grid timing.
    p = "ksi_lisims_lng"
    R.set(p, "first_export_year", 2032,
          f"Major Projects Office, Ksi Lisims LNG (modified 14 Sep 2026): {MPO_KSI} - Uniper "
          "agreement of 29 Jul 2026, 'with first deliveries expected to begin 2032'. "
          f"Corroborated by BC Government release of 27 May 2026 (SEFE agreement): {BC_SEFE} - "
          "'deliveries expected to begin by the early 2030s'. Replaces 2029 from GEM GGIT "
          f"(Sep 2025) LatestPlannedStartYear; GEM wiki still says 2029 at its 25 Sep 2026 edit: {GEM_KSI}")
    offtake_src = (f"Shell 2 mtpa: NGI, {NGI_KSI} | TotalEnergies 2 mtpa, 20 years: {TOTAL_KSI} | "
                   f"SEFE 1 mtpa, up to 20 years (27 May 2026): {BC_SEFE} | "
                   f"Uniper 2 mtpa, up to 20 years (29 Jul 2026): {MPO_KSI} | "
                   f"Santos 1 mtpa, 20 years (14 Sep 2026): {NRCAN_SANTOS}")
    R.set(p, "offtake_contracted_mtpa", 8.0, "Sum of five published agreements. " + offtake_src)
    R.set(p, "offtake_share_of_capacity", 0.667, "Derived: offtake_contracted_mtpa / capacity_mtpa")
    R.set(p, "offtake_counterparties",
          "Shell plc (2 mtpa); TotalEnergies SE (2 mtpa); Uniper SE (2 mtpa); SEFE (1 mtpa); Santos (1 mtpa)",
          offtake_src)
    R.set(p, "offtake_terms",
          "20 years (TotalEnergies, Santos); up to 20 years (SEFE, Uniper)", offtake_src)
    R.set(p, "export_term_note",
          "GL-346, 40 year term, issued 15 Mar 2023. Early-expiry clause: the term ends 10 years "
          "after issuance if LNG exports have not begun, so exports must start by about March "
          "2033. First deliveries are now expected in 2032.",
          "40-year term from the licence filing. End year = licence_issued_year + "
          "authorised_export_term_years, conservative (term may run from first export instead). | "
          f"GEM wiki, Ksi Lisims FLNG Terminal (edited 25 Sep 2026): {GEM_KSI} - 'early "
          "expiration clause where the term of the license ends 10 years after the date of "
          "issuance' if exports have not commenced. Licence text not opened (CER REGDOCS "
          "unreachable, 26 Sep 2026).")
    R.set(p, "grid_connection_status",
          "MOU with BC Hydro signed 20 Jan 2026 for up to 600 MW over a ~120 km Nass Valley "
          "Regional Transmission Line, conditional on the North Coast Transmission Line (NCTL) "
          "expansion. NCTL phase 2, west toward Terrace, starts construction in 2027 and is due "
          "to be complete by winter 2032.",
          f"Major Projects Office, Ksi Lisims LNG: {MPO_KSI} (600 MW, Nass Valley line) | "
          f"BC Government news release 2026ECS0002: {BC_HYDRO_KSI} | "
          f"BC Government, North Coast Transmission Line (Sep 2026): {BC_NCTL} (phase 2 complete by winter 2032)")
    R.set(p, "tier_reason",
          "Environmental approvals 15 Sep 2025; referred to the Major Projects Office 13 Nov 2025. "
          "Offtake agreements now cover about 8 of 12 mtpa (Shell, TotalEnergies, SEFE, Uniper, "
          "Santos). First deliveries expected in 2032 under the Uniper agreement. No FID as of "
          "27 Sep 2026.",
          f"Major Projects Office, Ksi Lisims LNG (modified 14 Sep 2026): {MPO_KSI} | "
          f"NRCan, third international offtake deal (15 Sep 2026): {NRCAN_SANTOS} | GEM wiki: {GEM_KSI}")
    R.add_sources(p,
                  f"BC Government, Ksi Lisims and SEFE agreement (27 May 2026): {BC_SEFE}",
                  f"NRCan, third international offtake deal (15 Sep 2026): {NRCAN_SANTOS}",
                  f"BC Government, North Coast Transmission Line: {BC_NCTL}")
    R.verified(p)

    # 3. Fermeuse: developer figure 10 mtpa; life derived from the developer's resource.
    p = "fermeuse_energy_flng"
    R.set(p, "capacity_mtpa", 10.0,
          f"The Narwhal, 19 Mar 2026: {NARWHAL_F} - Crown LNG chief executive Swapan Kataria: the "
          "plant would 'process and export up to 10 million tonnes of LNG per year'. A developer "
          "statement to journalists, not a filed or published nameplate. Adopted 27 Sep 2026 in "
          f"place of the 4.5 mtpa derivation. Same figure on GEM wiki: {GEM_FERM}")
    R.set(p, "capacity_basis",
          "Developer statement: up to 10 mtpa, from Crown LNG's chief executive in The Narwhal "
          "(19 Mar 2026). No project description or LNG plant application has been filed; only "
          "the marine base is approved. Feedgas would come through a 380 km subsea pipeline from "
          "the Jeanne d'Arc Basin (same source). HISTORY: 4.5 mtpa from 3 to 27 Sep 2026 and "
          "5.0 mtpa before that, both derived from the developer's 9.7 Tcf resource over a "
          "40-year life because no stated capacity had been found. That derivation is "
          "superseded; the resource claim now sets the operating life instead (see "
          "project_life_years).",
          f"The Narwhal, 19 Mar 2026: {NARWHAL_F}")
    R.set(p, "project_life_years", 18,
          "Derived, not stated. The developer's 9.7 Tcf (275 bcm) Jeanne d'Arc resource "
          "(Fermeuse Energy news release, 2 Sep 2025) divided by annual feedgas at the "
          f"developer's 10 mtpa (The Narwhal, 19 Mar 2026: {NARWHAL_F}): 10 mtpa x 1.38 bcm per "
          "mtpa x 1.1055 t feedgas per t LNG (fuel netted at the model's 0.29 / 2.75 factors, as "
          "in the superseded capacity derivation) = 15.26 bcm a year; 275 / 15.26 = 18.0 years at "
          "full output. Holding life to the resource keeps lifetime gas within the developer's "
          "own claim: 10 mtpa for 40 years would need about 610 bcm (21.5 Tcf). Model note: for "
          "pre-FID assets the FID delay sits inside this window, which shortens operating years "
          "(code decision pending).")
    R.set(p, "tier_reason",
          "Proposed floating LNG at the Fermeuse Marine Base. MOU with Hanwha Ocean Jan 2026. "
          "The developer has stated up to 10 mtpa (Mar 2026) but has filed no LNG plant "
          "proposal; only the marine base is approved.",
          f"The Narwhal, 19 Mar 2026: {NARWHAL_F} | GEM GGIT LNG Terminals (Sep 2025): {GEM_GGIT} | "
          f"GEM wiki: {GEM_FERM}")
    R.add_sources(p, f"The Narwhal, LNG in Newfoundland: lessons from Kitimat (19 Mar 2026): {NARWHAL_F}")
    R.verified(p)

    # 4. Cedar: latest permitted throughput 3.75 mtpa.
    p = "cedar_lng"
    R.set(p, "capacity_mtpa", 3.75,
          f"Cedar LNG application to amend its BC environmental assessment certificate (BC EAO, EPIC): {EAO_CEDAR} - "
          "liquefaction from 400 to 500 million standard cubic feet a day, from about 3.0 to "
          "3.75 mtpa, through efficiencies realised in detailed design with no change to "
          f"equipment. Amended certificate issued by the BC EAO in July 2026: Canadian Vanguard: {VANGUARD} "
          f"(9 Jul 2026); GIIGNL: {GIIGNL} (22 Jul 2026), which notes the change remains subject "
          "to approvals under the Impact Assessment Act and from the BC Energy Regulator. "
          f"Replaces 3.3 mtpa from the FID announcement (25 Jun 2024): {CEDAR_FID}")
    R.set(p, "capacity_basis",
          "Latest permitted design throughput: 3.75 mtpa (500 mmscfd) under the BC EAO's amended "
          "certificate of July 2026, from efficiencies with the same equipment, still subject to "
          "IAAC and BC Energy Regulator approval. Earlier figures: 3.3 mtpa in the developer's "
          "FID announcement (June 2024); 3.0 mtpa (400 mmscfd) from NRCan and the CER, the "
          "pre-amendment design basis.",
          f"BC EAO amendment application: {EAO_CEDAR} | Canadian Vanguard: {VANGUARD} | GIIGNL: {GIIGNL}")
    R.set(p, "offtake_share_of_capacity", 0.4, "Derived: offtake_contracted_mtpa / capacity_mtpa")
    R.set(p, "first_export_year", 2028,
          f"Pembina, 15 Dec 2025: {PEMBINA_DEC25} - 'expected in-service date in late 2028'. "
          f"Cedar LNG spring 2026 newsletter: {CEDAR_SPRING26} - FLNG vessel more than 50% "
          "complete, scheduled for delivery to Kitimat in 2028. Consistent with GEM GGIT "
          "(Sep 2025) LatestPlannedStartYear=2028.")
    R.set(p, "tier_reason",
          "Positive FID June 2024. FLNG vessel more than 50% complete at Samsung Heavy "
          "Industries (April 2026) and due at Kitimat in 2028; in service expected late 2028.",
          f"Cedar LNG spring 2026 newsletter: {CEDAR_SPRING26} | Pembina, 15 Dec 2025: {PEMBINA_DEC25}")
    R.add_sources(p,
                  f"BC EAO, Cedar LNG capacity amendment application: {EAO_CEDAR}",
                  f"Canadian Vanguard, Cedar LNG secures BC regulatory nod (Jul 2026): {VANGUARD}",
                  f"GIIGNL, Cedar LNG receives approval to increase project capacity (22 Jul 2026): {GIIGNL}",
                  f"Pembina, 2026 guidance and Cedar update (15 Dec 2025): {PEMBINA_DEC25}",
                  f"Cedar LNG spring 2026 newsletter: {CEDAR_SPRING26}")
    R.verified(p)

    # 5. Kino Aski: developer now names Western Canadian gas; drive claim; commercial news.
    p = "marinvest_baie_comeau"
    basin_src = (f"Kino Aski LNG website: {KINO_WEB} - '~1,000 KM' of 'new pipeline infrastructure "
                 f"extending Canadian gas networks to Baie-Comeau' | gasworld: {GASWORLD} - 'roughly "
                 "1,000km of new and existing pipeline infrastructure to transport natural gas from "
                 "Western Canada to Quebec' | Canada Investment Summit Prospectus (Invest in Canada, "
                 f"Sep 2026), p. 8: {PROSPECTUS} - 'connect Western Canadian natural gas supplies to "
                 "the Port of Baie-Comeau'")
    R.set(p, "feedgas_basin", "Western Canada Sedimentary Basin (developer-stated; producing area not named)",
          basin_src)
    R.set(p, "feedgas_basin_note",
          "The developer now names Western Canadian gas, carried over existing networks plus "
          "about 1,000 km of new pipeline to Baie-Comeau. The route is still undefined: three "
          "corridors are under study, including a new central route through Atikamekw "
          "Nitaskinan territory. No source names US supply, so the US Appalachian case in the "
          "Kino Aski feedgas sensitivity no longer has a basis (code change pending).",
          basin_src + f" | Québec Nouvelles, relaying Radio-Canada (Sep 2026): {QN} - third route")
    R.set(p, "liquefaction_drive_note",
          "previously classified as electric_planned. The developer plans up to three "
          "electric-drive FLNG units powered by about 1,700 MW of planned wind capacity. No "
          "contracted supply and no filed plan; electric drive is not assumed.",
          f"Kino Aski LNG website: {KINO_WEB}")
    R.set(p, "electrification_commitment",
          "Aspirational - about 1,700 MW of planned wind for up to three electric-drive FLNG "
          "units; no contracted supply. Electric drive is not assumed.",
          f"Kino Aski LNG website: {KINO_WEB}")
    R.set(p, "commercial_note",
          "No regulatory process has begun; the federal prospectus lists the project at "
          "'Pre-Application' with a capital cost of US$23B (Sep 2026), while Radio-Canada "
          "reporting cites about C$30bn. Kino Aski Inc. (Atikamekw-led) holds the majority and "
          "Marinvest Energy Canada a minority. Naftogaz signed a memorandum of understanding on "
          "11 Sep 2026 to assess buying Canadian LNG. The Anishnabe Nation of Lac-Simon says the "
          "project lacks community consent. Electric drive is claimed but not assumed.",
          f"Canada Investment Summit Prospectus, p. 8: {PROSPECTUS} | Québec Nouvelles: {QN} | "
          f"Kino Aski Inc. news release, 17 Aug 2026: {KINO_NEWSWIRE} | The Energy Mix, 16 Sep 2026: {ENERGYMIX_KINO}")
    R.set(p, "tier_reason",
          R.get(p, "tier_reason").rstrip()
          + " Listed as 'Pre-Application' in the federal Canada Investment Summit Prospectus (Sep 2026).",
          f"Kino Aski Inc. news release, 17 August 2026, {KINO_NEWSWIRE} | Canada Investment Summit Prospectus, p. 8: {PROSPECTUS}")
    R.add_sources(p,
                  f"Kino Aski LNG website: {KINO_WEB}",
                  f"gasworld, Kino Aski LNG targets North America-Europe energy corridor: {GASWORLD}",
                  f"Invest in Canada, Canada Investment Summit Prospectus (Sep 2026): {PROSPECTUS}",
                  f"The Energy Mix, Kino Aski signs tentative deal with Naftogaz (16 Sep 2026): {ENERGYMIX_KINO}",
                  f"Québec Nouvelles, third route studied (Sep 2026): {QN}")
    R.verified(p)

    # 6. Kanata: indicative terms with Hanwha.
    p = "kanata_lng"
    R.set(p, "tier_reason",
          "Announced 16 June 2026 via a non-binding MOU between Kanata Clean Power and Hanwha "
          "Ocean; estimated capital cost US$15.7 billion. On 14 Sep 2026 the parties agreed "
          "indicative terms: Hanwha Ocean's Energy Plant Unit would build, own and operate the "
          "facility as strategic partner and investor. A pre-FEED study follows a definitive "
          "agreement, full FEED is planned for 2027 and FID is targeted for early 2028. Final "
          "site selection is expected in Q4 2026. No regulatory filing has been made and the "
          "gas transport route is unspecified.",
          f"Hanwha Ocean news release, 16 June 2026, {HANWHA_KANATA} | LNG Industry, 14 Sep 2026: {LNGI_KANATA}")
    R.set(p, "owners",
          "Kanata Clean Power; Hanwha Ocean Energy Plant Unit (indicative terms, Sep 2026); a "
          "First Nation in the Prince Rupert region holds an exclusive right to acquire a 50/50 "
          "joint venture interest",
          f"LNG Industry, 14 Sep 2026: {LNGI_KANATA}")
    R.add_sources(p, f"LNG Industry, Kanata and Hanwha Ocean to launch pre-FEED (14 Sep 2026): {LNGI_KANATA}")
    R.verified(p)

    # 7. Woodfibre: in service end-2027; licence deadline; expansion note.
    p = "woodfibre_lng"
    R.set(p, "first_export_year", 2027,
          f"LNG Journal, 16 Jul 2026: {LNGJ_WF} - chief executive Luke Schauerte 'targets "
          "end-2027 for in-service'; project more than halfway complete. Woodfibre LNG "
          f"construction page: {WF_CONSTR} - 'substantial completion of the Project is expected "
          "by 2027'. Replaces 2028 from GEM GGIT (Sep 2025) under the 27 Sep 2026 rule that the "
          "developer's latest stated in-service year is used. An end-of-year start means little "
          "LNG in 2027 itself; the model's 40% first-year ramp overstates that year slightly.")
    R.set(p, "tier_reason",
          "Under construction and more than halfway complete (July 2026); in service targeted "
          "for end-2027. BC Hydro interconnection works scheduled 2026-2027.",
          f"LNG Journal, 16 Jul 2026: {LNGJ_WF} | GEM GGIT LNG Terminals (Sep 2025): {GEM_GGIT} | "
          "GEM wiki: https://www.gem.wiki/Woodfibre_LNG_Terminal")
    deadline = (" The CER extended the deadline for first exports under GL-340 from June 2027 "
                "to June 2030 (Aug 2025).")
    R.set(p, "export_term_note", R.get(p, "export_term_note").rstrip() + deadline,
          "40-year term from the licence filing. End year = licence_issued_year + "
          "authorised_export_term_years, conservative (term may run from first export instead). | "
          f"Energetic City, 29 Aug 2025: {ENERGETIC} (deadline extension)")
    R.set(p, "export_licence", R.get(p, "export_licence").rstrip() + deadline,
          f"Energetic City, 29 Aug 2025: {ENERGETIC}")
    R.set(p, "commercial_note",
          "100% of capacity contracted. A phased expansion from 2.1 to 2.5 mtpa (listed capex "
          "US$9.9B, 'Concept' stage) appears in the federal Canada Investment Summit Prospectus "
          "(Sep 2026). No filing located, so capacity stays at the 2.1 nameplate.",
          f"Compiled by us | Canada Investment Summit Prospectus, p. 10: {PROSPECTUS}")
    R.add_sources(p,
                  f"LNG Journal, Woodfibre on track to complete construction by late 2027 (16 Jul 2026): {LNGJ_WF}",
                  f"Woodfibre LNG, construction: {WF_CONSTR}",
                  f"Energetic City, export licence extension for Woodfibre LNG (29 Aug 2025): {ENERGETIC}",
                  f"Invest in Canada, Canada Investment Summit Prospectus (Sep 2026): {PROSPECTUS}")
    R.verified(p)

    # 8. Tilbury Phase 2: certificate and federal decision issued.
    p = "tilbury_phase_2"
    R.set(p, "tier_reason",
          "BC environmental assessment certificate and federal decision statement both issued "
          "21 Sep 2026 through a substituted assessment. Certified for up to 7,700 tonnes a day "
          "of new liquefaction and a 142,400 m3 tank, at about C$3.1bn with a six-year build. "
          "No FID. On the federal export project list under a 25 year export licence.",
          f"Impact Assessment Agency of Canada, 21 Sep 2026: {IAAC_TIL} | Surrey Now-Leader, "
          f"21 Sep 2026: {SURREY} | NRCan Canadian LNG projects: {NRCAN_LIST}")
    R.set(p, "capacity_basis",
          "GEM records Phase 2 as a 2.5 mtpa addition. The environmental assessment certificate "
          "(21 Sep 2026) covers up to 7,700 tonnes a day of new liquefaction, 2.8 Mt a year at "
          "365 days, so 2.5 mtpa is consistent with an addition rather than a site total. The "
          "earlier local reading of an addition of about 1.6 is set aside. GEM's 2.5 retained.",
          f"GEM GGIT LNG Terminals (Sep 2025): {GEM_GGIT} | Surrey Now-Leader, 21 Sep 2026: {SURREY}")
    R.add_sources(p,
                  f"Impact Assessment Agency of Canada, Tilbury Phase 2 decision (21 Sep 2026): {IAAC_TIL}",
                  f"Surrey Now-Leader, Tilbury Phase 2 EA certificate (21 Sep 2026): {SURREY}")
    R.verified(p)

    # 9. Tilbury Phase 1b: in service 2031.
    p = "tilbury_phase_1b"
    R.set(p, "first_export_year", 2031,
          f"BC Government news release 2026ECS0043 (Order in Council of 24 Jul 2026): {BC_T1B} - "
          "construction 'as early as mid-2027', 'in service as early as 2031'. Replaces 2028 "
          "from GEM GGIT (Sep 2025).")
    R.set(p, "tier_reason",
          "Proposed expansion, primarily marine bunkering and domestic supply; not listed "
          "separately by NRCan. An Order in Council of 24 Jul 2026 exempts it from needing a "
          "Certificate of Public Convenience and Necessity. More than C$2bn, construction from "
          "mid-2027, in service as early as 2031, with an equity option for the Musqueam Indian Band.",
          f"NRCan Canadian LNG projects: {NRCAN_LIST} | BC Government news release 2026ECS0043: {BC_T1B}")
    R.add_sources(p, f"BC Government, Tilbury Phase 1B Order in Council (24 Jul 2026): {BC_T1B}")
    R.verified(p)

    # 10. Tilbury Marine Jetty: proponent named.
    p = "tilbury_marine_jetty"
    jetty_src = (f"Tilbury Pacific Marine Jetty, About: {TILPAC} - 'Tilbury Jetty Limited "
                 "Partnership is the proponent of the Tilbury Pacific Marine Jetty Project, an "
                 "affiliate of FortisBC Holdings Inc.'")
    R.set(p, "proponent", "Tilbury Jetty Limited Partnership (affiliate of FortisBC Holdings Inc.)", jetty_src)
    R.set(p, "owners", "FortisBC Holdings Inc. (via Tilbury Jetty Limited Partnership)", jetty_src)
    R.set(p, "commercial_note",
          "Now named the Tilbury Pacific Marine Jetty Project. Would supply LNG to bunkering "
          "vessels in the Port of Vancouver. No service date announced.",
          f"Tilbury Pacific Marine Jetty, About: {TILPAC}")
    R.add_sources(p, f"Tilbury Pacific Marine Jetty, About: {TILPAC}")
    R.verified(p)

    # 11. Kitimat LNG: Haisla ownership; inactive to watch.
    p = "kitimat_t1t3"
    R.set(p, "proponent", "KM LNG Operating General Partnership (Bish LNG GP Ltd.)",
          f"BC EAO project record (EPIC): {EPIC_KIT}")
    R.set(p, "owners", "yáqʷa Development Corporation (Haisla Nation), acquired from Chevron Canada October 2025",
          f"yáqʷa Development Corporation, 17 Nov 2025: {YAQWA}")
    R.set(p, "status", "shelved",
          f"BC EAO project record (EPIC): {EPIC_KIT} - 'Post Decision - Care & Maintenance'. "
          f"The Chevron-Woodside project was cancelled (GEM GGIT, Sep 2025: {GEM_GGIT}); the "
          "certified site and permits survive under new ownership.")
    R.set(p, "tier", "watch", "Assigned by us - see tier_reason")
    R.set(p, "calc_group", "watch", "Assigned: watch rows are excluded from all totals")
    R.set(p, "chain", "export",
          "Assigned by us: former export proposal. Needed because the row carries a capacity; "
          "watch rows are excluded from all totals.")
    R.set(p, "export_licence",
          "Export licences transferred to yáqʷa with the site (October 2025); licence numbers and terms not verified.",
          f"yáqʷa Development Corporation, 17 Nov 2025: {YAQWA}")
    R.set(p, "tier_reason",
          "Chevron and Woodside's Kitimat LNG was cancelled (GEM, 2021). In October 2025 yáqʷa "
          "Development Corporation, owned by the Haisla Nation, acquired from Chevron 'all "
          "relevant permits and licenses associated with Kitimat LNG, including Environmental "
          "Assessment Certificate, Export Licenses', with site remediation funding. The BC EAO "
          "lists the project as 'Post Decision - Care & Maintenance'. No development plan, "
          "capacity or partner has been announced. Moved from inactive to watch on 27 Sep 2026: "
          "a fully permitted site with no live proposal. Capacity is the cancelled project's GEM "
          "figure and enters no total.",
          f"yáqʷa Development Corporation, 17 Nov 2025: {YAQWA} | BC EAO project record (EPIC): {EPIC_KIT} | "
          f"GEM wiki: {GEM_KIT}")
    R.add_sources(p,
                  f"yáqʷa Development Corporation, acquisition of Kitimat LNG assets (17 Nov 2025): {YAQWA}",
                  f"BC EAO project record, Kitimat LNG (EPIC): {EPIC_KIT}")
    R.verified(p)

    # 13. Placentia Bay: EA registration still active.
    p = "placentia_bay_flng"
    R.set(p, "proponent", "LNG Newfoundland and Labrador Limited",
          f"Newfoundland and Labrador environmental assessment projects list: {NL_EA} - registration 2177")
    R.set(p, "tier_reason",
          "Shelved. Pre-FID. Its provincial environmental assessment registration (no. 2177, "
          "LNG Newfoundland and Labrador Limited, 23 Nov 2021) is still listed as Active with no "
          "decision (NL registry, checked 27 Sep 2026).",
          f"GEM GGIT LNG Terminals (Sep 2025): {GEM_GGIT} | GEM wiki: https://www.gem.wiki/Placentia_Bay_FLNG_Terminal | "
          f"NL environmental assessment projects list: {NL_EA}")
    R.add_sources(p, f"Newfoundland and Labrador environmental assessment projects list: {NL_EA}")
    R.verified(p)

    # 14. Saint John: capacity conflict recorded.
    p = "saint_john_import_facility"
    R.set(p, "capacity_basis",
          "nameplate (GEM): 7.5 mtpa of import (regasification) capacity. NRCan's LNG project "
          "list gives 6.5 mtpa (page modified 7 Jan 2025). Conflict recorded, not resolved; the "
          "model uses terminal-specific throughput for this asset, so the figure barely affects results.",
          f"GEM GGIT LNG Terminals (Sep 2025): {GEM_GGIT} | NRCan Canadian LNG projects: {NRCAN_LIST}")
    R.add_sources(p, f"NRCan Canadian LNG projects: {NRCAN_LIST}")
    R.verified(p)

    # 15. Tamaska: owner's current name.
    p = "tamaska_fort_nelson_lng"
    cryo_src = (f"Cryopeak Energy Solutions website: {CRYOPEAK} - current company name; 'owns and "
                "operates three advanced LNG production facilities throughout Western Canada'. "
                f"NRCan listed the owner as Cryopeak LNG Solutions: {NRCAN_LIST}")
    R.set(p, "proponent", "Cryopeak Energy Solutions", cryo_src)
    R.set(p, "owners", "Cryopeak Energy Solutions", cryo_src)
    R.add_sources(p, f"Cryopeak Energy Solutions: {CRYOPEAK}")
    R.verified(p)


def update_supporting_infrastructure(R: Register):
    ws = R.wb["Supporting Infrastructure"]
    cols = {c.value: c.column for c in ws[1] if c.value}
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, cols["pipeline_id"]).value == "coastal_gaslink":
            cap = 14.0 + 3.75
            demand = cap * 1.38  # bcm/y of gas-equivalent LNG output, as in the existing rows
            headroom = ws.cell(r, cols["capacity_bcm_y"]).value - demand
            ws.cell(r, cols["terminal_capacity_mtpa"]).value = cap
            ws.cell(r, cols["feedgas_demand_bcm_y"]).value = round(demand + 1e-9, 2)
            ws.cell(r, cols["headroom_bcm_y"]).value = round(headroom - 1e-9, 2)
            ws.cell(r, cols["capacity_verdict"]).value = f"SHORTFALL of {abs(round(headroom - 1e-9, 2)):.2f} bcm/y"
            src = ws.cell(r, cols["sources"]).value or ""
            if EAO_CEDAR not in src:
                ws.cell(r, cols["sources"]).value = (
                    src + f" | Cedar at 3.75 mtpa from 27 Sep 2026 (BC EAO amendment): {EAO_CEDAR}")
            R.log.append(("SupportingInfrastructure:coastal_gaslink", "terminal/demand/headroom",
                          "17.3 / 23.87 / -2.41", f"{cap} / {round(demand + 1e-9, 2)} / {round(headroom - 1e-9, 2)}"))


def update_data_gaps(R: Register):
    ws = R.wb["Data Gaps"]
    hdr = {c.value: c.column for c in ws[1] if c.value}

    def row_for(field: str, rows_affected_start: str) -> int:
        for r in range(2, ws.max_row + 1):
            if (ws.cell(r, hdr["field"]).value == field
                    and str(ws.cell(r, hdr["rows_affected"]).value).startswith(rows_affected_start)):
                return r
        raise KeyError((field, rows_affected_start))

    def put(r: int, **vals):
        for k, v in vals.items():
            old = ws.cell(r, hdr[k]).value
            ws.cell(r, hdr[k]).value = v
            R.log.append((f"DataGaps:row{r}", k, old, v))

    put(row_for("capacity_mtpa", "Fermeuse"),
        current_state="Developer statement 10 mtpa; life derived 18 years",
        why=("Crown LNG's chief executive stated up to 10 mtpa (The Narwhal, 19 Mar 2026: "
             f"{NARWHAL_F}). No filing states capacity or life. The 18-year life is derived from "
             "the developer's 9.7 Tcf resource at 10 mtpa, net of liquefaction fuel. Supersedes "
             "the 4.5 mtpa derivation of 3 Sep 2026."),
        what_would_fill_it="A filed project description stating liquefaction capacity and production profile.",
        priority="high - a developer statement, not a filing; the derived life sets the lifetime total")
    put(row_for("capacity_mtpa", "Cedar"),
        current_state="UPDATED - 3.75 permitted",
        why=("The BC EAO's amended certificate (July 2026) permits 500 mmscfd, 3.75 mtpa, from "
             f"the same equipment ({EAO_CEDAR}; {GIIGNL}). The FID announcement said 3.3; NRCan "
             "and the CER give 3.0, the pre-amendment basis. Still subject to IAAC and BC Energy "
             "Regulator approval."),
        what_would_fill_it="IAAC decision statement amendment and BC Energy Regulator permit",
        priority="low - 0.45 mtpa")
    put(row_for("capacity_mtpa", "Tilbury Phase 2"),
        current_state="RESOLVED - 2.5 addition retained",
        why=("The 21 Sep 2026 certificate covers up to 7,700 t/d of new liquefaction (2.8 Mt a "
             f"year at 365 days: {SURREY}), consistent with GEM's 2.5 mtpa addition and not with "
             "the ~1.6 implied by earlier local reporting."),
        what_would_fill_it="Closed",
        priority="closed")
    put(row_for("first_export_year", "Woodfibre"),
        current_state="RESOLVED - 2027 used",
        why=("The developer's latest statement targets in service at end-2027 (LNG Journal, "
             f"16 Jul 2026: {LNGJ_WF}). GEM (Sep 2025) said 2028; Norton Rose and the CER 2027. "
             "The latest developer year is used (27 Sep 2026)."),
        what_would_fill_it="Closed; revisit if first cargo slips into 2028",
        priority="closed")
    put(row_for("proponent", "Tilbury Marine Jetty"),
        current_state="RESOLVED - Tilbury Jetty Limited Partnership",
        why=(f"The project site ({TILPAC}) names Tilbury Jetty Limited Partnership, a FortisBC "
             "Holdings affiliate, as proponent. Whether WesPac's GL-310 relates to this project "
             "is still unconfirmed."),
        what_would_fill_it="CER REGDOCS for the GL-310 link",
        priority="low")
    put(row_for("feedgas_basin", "Kino Aski"),
        current_state="RESOLVED at basin level - Western Canada (developer-stated)",
        why=("The developer names Western Canadian gas carried over existing networks plus "
             f"about 1,000 km of new pipeline ({KINO_WEB}; {GASWORLD}; federal prospectus p. 8, "
             f"{PROSPECTUS}). No source states US supply."),
        what_would_fill_it="A filed route naming the producing area and pipeline corridor.",
        priority="medium - the Canadian territorial attribution is settled; the pipeline distance is not")
    r = row_for("pipeline route distance", "Kino Aski")
    put(r, why=(ws.cell(r, hdr["why"]).value.rstrip()
                + " Update Sep 2026: the developer describes about 1,000 km of new pipeline added "
                "to existing Canadian networks, and three corridors are under study "
                f"({KINO_WEB}; {QN}), so a route distance still cannot be computed."))


def update_readme_and_definitions(R: Register):
    ws = R.wb["README"]

    def row_of(label: str) -> int:
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == label:
                return r
        raise KeyError(label)

    def put(label: str, text: str, new_label: str | None = None):
        r = row_of(label)
        R.log.append((f"README:{label}", "B", ws.cell(r, 2).value, text))
        ws.cell(r, 2).value = text
        if new_label:
            ws.cell(r, 1).value = new_label

    put("Build", "2026-09-27 (register update, see REGISTER UPDATE 27 SEP 2026 below). Previous build 2026-08-20 10:00 UTC.")
    put("Scope", ws.cell(row_of("Scope"), 2).value.replace(
        "Discovery LNG is in GEM but was moved from inactive to early_proposed pending verification.",
        "Discovery LNG is in GEM as cancelled and the register matches (1 Sep 2026)."))
    put("early_proposed",
        "Announced, but no completed regulatory filing. Groups with advanced_proposed as proposed "
        "in the calculation, but the tier is retained so the difference in maturity stays visible. "
        "Capacity figures here are weaker: Kino Aski's 15 mtpa is a developer press release and "
        "Fermeuse's 10 mtpa a developer statement to journalists, neither a filed figure, and "
        "Summit Lake's 2.7 is an upper bound from a suspended assessment.")
    put("watch", ws.cell(row_of("watch"), 2).value.rstrip()
        + " Kitimat LNG is in watch too: a permitted site held by the Haisla Nation with no live proposal.")
    put("The standard is nameplate",
        "capacity_mtpa records the most recent design capacity published by the developer or "
        "permitted by a regulator. Where a regulator has permitted more than the original "
        "nameplate from the same equipment, the permitted figure is used (decision of 27 Sep "
        "2026: use the latest available data).",
        new_label="The standard is the latest design figure")
    put("Alternatives are recorded",
        "Where sources give different figures, the latest design figure is used and the others "
        "are stated in capacity_basis. Cedar LNG: the FID announcement stated 3.3 mtpa and NRCan "
        "and the CER publish 3.0, but the BC EAO's amended certificate of July 2026 permits 3.75 "
        "mtpa from the same equipment. The register carries 3.75.")
    put("Not every figure is a nameplate",
        "Kino Aski's 15 mtpa is from the developer's 17 August 2026 press release rather than a "
        "filed project description; the earlier 10 mtpa journalist statement is replaced. Its "
        "feedgas route is undefined, though the developer now names Western Canadian gas. "
        "Fermeuse's 10 mtpa is a developer statement to The Narwhal (March 2026). Summit Lake's "
        "2.7 is an upper bound from its impact assessment. All are flagged in capacity_basis.")
    put("Four projects hold 40 year licences", ws.cell(row_of("Four projects hold 40 year licences"), 2).value.rstrip()
        + " Ksi Lisims' licence also ends early if exports have not begun within 10 years of issue (GEM).")
    put("Two caveats",
        "A licence term is not an operating life: licences are renewable and Tilbury has run "
        "since 1971. Since Bill C-15 received royal assent on 26 March 2026, LNG export licences "
        f"can run for up to 50 years (Major Projects Office: {MPO_P2}), so 40 is not a ceiling. "
        "The 30/40/50 sensitivity lives in the Data Inputs file.")
    put("None of it serves Canadians", ws.cell(row_of("None of it serves Canadians"), 2).value.replace(
        "All 45.4 mtpa", "All 45.85 mtpa"))
    put("Coastal GasLink is over-subscribed",
        "It carries 21.46 bcm/y and already serves 24.50 from LNG Canada Phase 1 plus Cedar (3.75 "
        "mtpa since the July 2026 certificate amendment), whose 8 km connector draws on it. Phase "
        "2's 14 mtpa, 31 per cent of live export capacity, depends on the proposed second line "
        "agreed only in principle in March 2026.")
    put("Aggregate headroom is misleading", ws.cell(row_of("Aggregate headroom is misleading"), 2).value.replace(
        "against 62.65 of demand", "against 63.28 of demand"))
    put("Fermeuse capacity is derived",
        "Crown LNG's chief executive told The Narwhal (19 Mar 2026) the plant would export up to "
        "10 mtpa, and the register carries 10. No filing states it. The operating life is "
        "derived instead: the developer's 9.7 Tcf Jeanne d'Arc resource lasts about 18 years at "
        "10 mtpa, net of liquefaction fuel, recorded in project_life_years. The earlier 4.5 mtpa "
        "derivation over a 40-year life is superseded and kept as history in capacity_basis.",
        new_label="Fermeuse capacity is a developer statement")

    # Change log block at the foot of the README sheet.
    style_a, style_b = ws.cell(55, 1), ws.cell(55, 2)
    head_style = ws.cell(51, 1)
    start = ws.max_row + 2
    entries = [
        ("REGISTER UPDATE 27 SEP 2026", None),
        ("Not yet run", "The model has not been re-run on this register. Every result in the repo README and Outputs is still the 3 September 2026 lock."),
        ("LNG Canada Phase 2", f"No FID yet. Status text refreshed; first LNG 2030 reconfirmed (GEM wiki, 25 Sep 2026: {GEM_LNGC})."),
        ("Ksi Lisims", f"first_export_year 2029 to 2032 (Uniper agreement, {MPO_KSI}); offtake 4 to 8 mtpa; licence early-expiry clause; grid timing."),
        ("Fermeuse", f"capacity 4.5 to 10 mtpa (developer statement, {NARWHAL_F}); project_life_years 18, derived from the developer's resource."),
        ("Cedar", f"capacity 3.3 to 3.75 mtpa (BC EAO amended certificate, July 2026: {EAO_CEDAR}). Coastal GasLink demand and headroom updated to match."),
        ("Kino Aski", f"feedgas basin Western Canada (developer-stated: {KINO_WEB}); drive, cost and commercial notes."),
        ("Kanata", f"Indicative terms with Hanwha, pre-FEED, FID targeted early 2028 ({LNGI_KANATA})."),
        ("Woodfibre", f"first_export_year 2028 to 2027 (developer targets end-2027 in service: {LNGJ_WF}); GL-340 deadline June 2030; expansion noted, capacity unchanged."),
        ("Tilbury Phase 2", f"Certificate and federal decision 21 Sep 2026 ({IAAC_TIL}); 2.5 mtpa confirmed as an addition."),
        ("Tilbury Phase 1b", f"first_export_year 2028 to 2031 ({BC_T1B})."),
        ("Tilbury Marine Jetty", f"Proponent named: Tilbury Jetty Limited Partnership ({TILPAC})."),
        ("Kitimat LNG", f"Owner yáqʷa Development Corporation (Haisla Nation) ({YAQWA}); inactive to watch."),
        ("Placentia Bay", f"EA registration 2177 still Active ({NL_EA})."),
        ("Saint John", f"NRCan's 6.5 mtpa import capacity recorded against GEM's 7.5 ({NRCAN_LIST})."),
        ("Tamaska", f"Owner's current name: Cryopeak Energy Solutions ({CRYOPEAK})."),
        ("Not changed", "Kitsault: the BusinessWire releases on a multi-commodity port could not be opened (404), so no change was made."),
    ]
    for i, (a, b) in enumerate(entries):
        r = start + i
        ca, cb = ws.cell(r, 1), ws.cell(r, 2)
        ca.value, cb.value = a, b
        src_a = head_style if b is None else style_a
        ca.font, ca.alignment, ca.fill, ca.border = copy(src_a.font), copy(src_a.alignment), copy(src_a.fill), copy(src_a.border)
        cb.font, cb.alignment, cb.fill, cb.border = copy(style_b.font), copy(style_b.alignment), copy(style_b.fill), copy(style_b.border)

    # Field definitions that the Cedar and Fermeuse changes make stale.
    fd = R.wb["Field Definitions"]
    for r in range(2, fd.max_row + 1):
        name = fd.cell(r, 1).value
        if name == "capacity_mtpa":
            fd.cell(r, 3).value = ("Design export capacity. The single most important input to every "
                                   "headline number. Convention: the latest design figure published by "
                                   "the developer or permitted by a regulator, with other figures "
                                   "recorded in capacity_basis where they differ.")
        if name == "project_life_years":
            fd.cell(r, 3).value = ("Operating life where the developer states one (Summit Lake PG LNG, "
                                   "30 years) or where it follows from the developer's own statements "
                                   "(Fermeuse, 18 years: stated resource at stated capacity, derivation "
                                   "in the source cell). There is deliberately no blanket default.")


def main(path: str):
    R = Register(path)
    update_asset_rows(R)
    update_supporting_infrastructure(R)
    update_data_gaps(R)
    update_readme_and_definitions(R)
    R.wb.save(path)
    for pid, field, old, new in R.log:
        o = str(old)[:70].replace("\n", " ")
        n = str(new)[:70].replace("\n", " ")
        print(f"{pid:42s} {field:28s} {o!s:72s} -> {n}")


if __name__ == "__main__":
    main(sys.argv[1])
